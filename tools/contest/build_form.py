"""제출 양식 4~7쪽(서식3 · 서식4)을 양식과 같은 모양의 Word 파일로 만든다. 한글(HWP)에서 열어 칸별로 옮겨 붙이면 된다.

실행: python3 tools/contest/build_form.py  →  docs/공모전_제출_서식3_4.docx (+ 같은 이름 .md)
그림: ui-shots 브랜치의 최신 실제 앱 캡처(ui_v48)에서 5장을 가져와 한 줄로 붙인다.
"""
import io
import os
import subprocess

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

import form_content as C

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
OUT = os.path.join(ROOT, 'docs', '공모전_제출_서식3_4')
SHOTS = 'ui_v48'
FONT = '맑은 고딕'
NAVY = RGBColor(0x00, 0x00, 0x80)
BLUE = RGBColor(0x00, 0x00, 0xFF)
HEAD_BG, LABEL_BG, NOTE_BG = 'DCE3F5', 'E6E6E6', 'E6E6E6'
PRETENDARD = os.path.join(ROOT, 'My project', 'Assets', 'Resources', 'KUCA', 'Fonts')


def shade(cell, hex_color):
    pr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hex_color)
    pr.append(sh)


def borders(table, size=8):
    tbl = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(size)); e.set(qn('w:color'), '000000')
        b.append(e)
    tbl.append(b)


def run(p, text, size=10, bold=False, color=None):
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if color:
        r.font.color.rgb = color
    return r


def put(cell, text, size=10, bold=False, color=None, center=False, first=True):
    """칸에 글을 넣는다. 줄바꿈마다 문단을 나눈다."""
    for i, line in enumerate(text.split('\n')):
        p = cell.paragraphs[0] if (first and i == 0) else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, line, size, bold, color)


def widths(table, mm):
    """칸 너비 고정 (LibreOffice · 한글은 gridCol 과 fixed 배치를 본다)"""
    table.autofit = False
    pr = table._tbl.tblPr
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); pr.append(layout)
    for col, w in zip(table._tbl.tblGrid.findall(qn('w:gridCol')), mm):
        col.set(qn('w:w'), str(Mm(w).twips))
    for row in table.rows:
        col = 0
        for tc in row._tr.tc_lst:
            n = tc.grid_span
            tc.width = Mm(sum(mm[col:col + n]))
            col += n


def form_bar(doc, no, title):
    """양식 머리: [서식N] 남색 칸 + 제목 칸"""
    t = doc.add_table(rows=1, cols=2); borders(t, 12)
    widths(t, [25, 145])
    shade(t.cell(0, 0), '000080')
    put(t.cell(0, 0), f'서식{no}', 14, True, RGBColor(0xFF, 0xFF, 0xFF), center=True)
    put(t.cell(0, 1), title, 14, True)
    doc.add_paragraph()


def section(doc, title, note, body=None):
    """회색 제목 줄 + 파란 안내 + 본문 칸"""
    t = doc.add_table(rows=2 if body is not None else 1, cols=1); borders(t)
    widths(t, [170])
    shade(t.cell(0, 0), NOTE_BG)
    put(t.cell(0, 0), title, 10.5, True)
    put(t.cell(0, 0), note, 9, color=BLUE, first=False)
    if body is not None:
        put(t.cell(1, 0), body, 10)
    return t


def figure_png():
    """실제 앱 캡처 5장을 한 줄로 (캡션 포함)"""
    W, cap = 360, 64
    font = ImageFont.truetype(os.path.join(PRETENDARD, 'Pretendard-Bold.otf'), 30)
    tiles = []
    for name, caption in C.S2_FIGURE:
        raw = subprocess.run(['git', '-C', ROOT, 'show', f'origin/ui-shots:{SHOTS}/{name}'], capture_output=True, check=True).stdout
        im = Image.open(io.BytesIO(raw)).convert('RGB')
        im = im.resize((W, int(im.size[1] * W / im.size[0])))
        tiles.append((im, caption))
    H = max(t.size[1] for t, _ in tiles)
    gap = 16
    out = Image.new('RGB', (len(tiles) * W + (len(tiles) - 1) * gap, H + cap), 'white')
    d = ImageDraw.Draw(out)
    for i, (im, caption) in enumerate(tiles):
        x = i * (W + gap)
        out.paste(im, (x, 0))
        d.rectangle([x, 0, x + W - 1, H - 1], outline=(200, 200, 200), width=2)
        tw = d.textlength(caption, font=font)
        d.text((x + (W - tw) / 2, H + 14), caption, fill=(20, 20, 20), font=font)
    buf = io.BytesIO(); out.save(buf, 'PNG'); buf.seek(0)
    return buf


def form3(doc):
    form_bar(doc, 3, '프로젝트 소개서')
    t = doc.add_table(rows=7, cols=4); borders(t)
    widths(t, [32, 54, 30, 54])
    t.cell(0, 0).merge(t.cell(0, 3)); shade(t.cell(0, 0), LABEL_BG)
    put(t.cell(0, 0), '프로젝트 소개서', 14, True, center=True)
    labels = ['부   문', '팀   명', '아이디어명(주제)', '작품 형태', '작품 제출 파일 목록', '작품 요약 소개\n(1,000자 이내)']
    for i, label in enumerate(labels, start=1):
        shade(t.cell(i, 0), LABEL_BG)
        put(t.cell(i, 0), label, 10, True, center=True)
        t.cell(i, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    # 부문
    t.cell(1, 1).merge(t.cell(1, 3))
    put(t.cell(1, 1), '☑ AI 문제해결 솔루션        □ AI 창작 콘텐츠', 10.5)
    # 팀명 · 팀장
    shade(t.cell(2, 2), LABEL_BG); put(t.cell(2, 2), '팀   장', 10, True, center=True)
    put(t.cell(2, 1), C.TEAM, 10.5); put(t.cell(2, 3), C.LEADER, 10.5, color=BLUE)
    for r in (3, 4, 5, 6):
        t.cell(r, 1).merge(t.cell(r, 3))
    put(t.cell(3, 1), C.TITLE, 10.5, True)
    put(t.cell(4, 1), '  '.join(('☑ ' if f == C.FORM_TYPE else '□ ') + f for f in C.FORM_TYPES), 10)
    put(t.cell(5, 1), '\n'.join(C.FILES), 9.5)
    put(t.cell(6, 1), C.SUMMARY, 9.5)
    doc.add_paragraph()

    section(doc, '1. 문제 인식과 주제 적합성 (1,000자 이내)',
            '※ 해결하려는 캠퍼스·일상·사회 문제 또는 창작 주제 기술 (주제 선정 배경, 필요성, 목적/목표, 기대효과 등)', C.S1)
    doc.add_paragraph()

    doc.add_page_break()
    t2 = section(doc, '2. 핵심 아이디어와 사용자 경험', '※ 주요 기능/스토리/화면·체험 순서. 필요 시 이미지·도표 삽입', '')
    body = t2.cell(1, 0)
    put(body, C.S2_INTRO, 10)
    body.add_paragraph().add_run().add_picture(figure_png(), width=Mm(164))
    put(body, '그림. 실제 앱 화면 (Unity 빌드에서 자동 캡처, 10/27)', 8.5, color=RGBColor(0x55, 0x55, 0x55), center=True, first=False)
    put(body, '체험 순서', 10, True, first=False)
    flow = body.add_table(rows=len(C.S2_FLOW), cols=2); borders(flow, 4)
    for i, (k, v) in enumerate(C.S2_FLOW):
        shade(flow.cell(i, 0), HEAD_BG)
        put(flow.cell(i, 0), k, 9.5, True, center=True)
        put(flow.cell(i, 1), v, 9)
    widths(flow, [20, 142])
    put(body, 'AI 가 하는 일과 사람이 정한 틀', 10, True, first=False)
    ai = body.add_table(rows=len(C.S2_AI) + 1, cols=3); borders(ai, 4)
    for j, h in enumerate(['기능', 'AI 가 하는 일', '사람이 정한 틀 · AI 실패 시']):
        shade(ai.cell(0, j), HEAD_BG); put(ai.cell(0, j), h, 9.5, True, center=True)
    for i, row in enumerate(C.S2_AI, start=1):
        for j, v in enumerate(row):
            put(ai.cell(i, j), v, 9, bold=(j == 0))
    widths(ai, [26, 66, 70])
    doc.add_page_break()
    section(doc, '3. 창의성·차별성 (500자 이내)', '※ 기존 사례와 비교한 독창적 요소, 단순 생성과 구별되는 사람의 기획·편집', C.S3)
    doc.add_paragraph()
    section(doc, '4. 실용성·완성도·확장성 (1,000자 이내)', '※ 현재 구현 수준, 작동 범위, 한계, 향후 고도화, 예상 효과', C.S4)


def form4(doc):
    doc.add_page_break()
    form_bar(doc, 4, 'AI 활용 과정 설명서')
    p = doc.add_paragraph(); run(p, '□ AI 활용 과정 및 참가자 기여 내역', 11, True)
    p = doc.add_paragraph(); run(p, '※ 심사위원이 AI 활용의 깊이와 참가자의 실질적 기여를 확인할 수 있게 작성', 9, color=BLUE)
    p = doc.add_paragraph(); run(p, '※ 증빙 원칙 : 전체 대화기록을 무조건 제출하기보다 핵심 프롬프트·버전 변화·검증 결과를 선별해 첨부하고, 개인정보·API 키·비공개 정보는 반드시 제거합니다.', 9, color=BLUE)

    t = doc.add_table(rows=1 + 2 * len(C.STAGES), cols=3); borders(t)
    t.cell(0, 1).merge(t.cell(0, 2))
    for j, h in ((0, '활용 단계'), (1, '내용')):
        shade(t.cell(0, j), HEAD_BG); put(t.cell(0, j), h, 10, True, center=True)
    for i, (stage, model, prompt) in enumerate(C.STAGES):
        r = 1 + 2 * i
        stage_cell = t.cell(r, 0).merge(t.cell(r + 1, 0))
        shade(stage_cell, LABEL_BG)
        put(stage_cell, stage, 10, True, center=True)
        stage_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for rr, label, text in ((r, 'AI 모델', model), (r + 1, '주요 프롬프트', prompt)):
            put(t.cell(rr, 1), label, 9.5, center=True)
            t.cell(rr, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            put(t.cell(rr, 2), text, 8.5)
    widths(t, [18, 24, 128])
    doc.add_paragraph()

    p = doc.add_paragraph(); run(p, '□ 출처·권리·라이선스 목록 ', 11, True); run(p, '(*해당 사항 있는 경우, 기재 요망)', 9, color=BLUE)
    rows = sum(len(v) for v in C.LICENSES.values())
    t = doc.add_table(rows=rows + 1, cols=5); borders(t)
    for j, h in enumerate(['자료 유형', '자료명·출처', '저작권자·제공처', '이용 조건', '활용 내용']):
        shade(t.cell(0, j), HEAD_BG); put(t.cell(0, j), h, 9.5, True, center=True)
    r = 1
    for kind, items in C.LICENSES.items():
        kc = t.cell(r, 0).merge(t.cell(r + len(items) - 1, 0))
        shade(kc, LABEL_BG); put(kc, kind, 9.5, True, center=True)
        kc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for item in items:
            for j, v in enumerate(item, start=1):
                put(t.cell(r, j), v, 8, color=BLUE if v.startswith('※') else None)
            r += 1
    widths(t, [20, 46, 36, 34, 34])


def markdown():
    """같은 내용을 붙여넣기용 텍스트로"""
    L = ['# 2026 KHU AI Challenge 제출 양식 4~7쪽 작성본 (서식3 · 서식4)', '',
         '> `tools/contest/build_form.py` 로 만든 Word 파일(`공모전_제출_서식3_4.docx`)과 같은 내용입니다. ※ 표시는 팀이 확인한 뒤 고칠 칸입니다.', '',
         '## 서식3 프로젝트 소개서', '',
         f'- 부문: ☑ {C.DIVISION}', f'- 팀명: {C.TEAM} · 팀장: {C.LEADER}', f'- 아이디어명: {C.TITLE}', f'- 작품 형태: ☑ {C.FORM_TYPE}',
         '- 작품 제출 파일 목록:'] + [f'  - {f}' for f in C.FILES] + ['']
    for name, text, limit in C.LIMITS[:1]:
        L += [f'### 작품 요약 소개 ({len(text)} / {limit}자)', '', text, '']
    L += [f'### 1. 문제 인식과 주제 적합성 ({len(C.S1)} / 1000자)', '', C.S1, '',
          '### 2. 핵심 아이디어와 사용자 경험', '', C.S2_INTRO, '', '| 단계 | 체험 |', '| --- | --- |']
    L += [f'| {k} | {v} |' for k, v in C.S2_FLOW]
    L += ['', '| 기능 | AI 가 하는 일 | 사람이 정한 틀 · AI 실패 시 |', '| --- | --- | --- |']
    L += [f'| {a} | {b} | {c} |' for a, b, c in C.S2_AI]
    L += ['', f'### 3. 창의성·차별성 ({len(C.S3)} / 500자)', '', C.S3, '',
          f'### 4. 실용성·완성도·확장성 ({len(C.S4)} / 1000자)', '', C.S4, '', '## 서식4 AI 활용 과정 설명서', '']
    for stage, model, prompt in C.STAGES:
        L += [f'### {stage}', '', f'- **AI 모델:** {model}', '- **주요 프롬프트:**', '', prompt, '']
    L += ['### 출처·권리·라이선스 목록', '', '| 자료 유형 | 자료명·출처 | 저작권자·제공처 | 이용 조건 | 활용 내용 |', '| --- | --- | --- | --- | --- |']
    for kind, items in C.LICENSES.items():
        L += [f'| {kind} | ' + ' | '.join(i) + ' |' for i in items]
    return '\n'.join(L).replace('\n→', '\n\n→') + '\n'


def main():
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = Mm(20)
        s.top_margin = s.bottom_margin = Mm(15)
    st = doc.styles['Normal']
    st.font.name = FONT; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    form3(doc)
    form4(doc)
    doc.save(OUT + '.docx')
    with open(OUT + '.md', 'w', encoding='utf-8') as f:
        f.write(markdown())
    print(OUT + '.docx')


if __name__ == '__main__':
    main()
